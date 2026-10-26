from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

from exam_seating.exporters import export_excel, export_pdf, unique_sheet_name
from exam_seating.models import SeatingSummary
from exam_seating.sirt_exporters import export_sirt_attendance_sheets_excel, export_sirt_debarred_excel

SUMMARY = SeatingSummary(1, 1, 0, 1, 0, 2, 2, 2)
NOT_ELIGIBLE = pd.DataFrame(columns=["Roll Number", "Name", "Attendance Percentage"])


def _room(roll: str, name: str) -> pd.DataFrame:
    return pd.DataFrame(
        {"Seat Number": [1], "Roll Number": [roll], "Name": [name], "Attendance Percentage": [80.0]}
    )


def test_unique_sheet_name_handles_clashes_case_and_length():
    used: set[str] = set()
    assert unique_sheet_name("Room_A/1", used) == "Room_A1"
    assert unique_sheet_name("room_a1", used) == "room_a1_2"  # Excel ignores case
    long = "x" * 40
    first, second = unique_sheet_name(long, used), unique_sheet_name(long, used)
    assert len(first) == 31 and len(second) == 31 and first != second


def test_rooms_with_clashing_names_do_not_crash_the_workbook(tmp_path: Path):
    rooms = {"A/1": _room("1", "A"), "A1": _room("2", "B")}
    path = export_excel(rooms, NOT_ELIGIBLE, SUMMARY, tmp_path)
    names = openpyxl.load_workbook(path).sheetnames
    assert len(names) == len({n.lower() for n in names})  # no duplicates
    assert {"Room_A1", "Room_A1_2"} <= set(names)


def test_names_are_never_turned_into_formulas_or_links(tmp_path: Path):
    rooms = {"X": _room("1", "=1+1")}
    rooms["Y"] = _room("2", "http://example.com")
    path = export_excel(rooms, NOT_ELIGIBLE, SUMMARY, tmp_path)
    book = openpyxl.load_workbook(path)
    formula_cell, link_cell = book["Room_X"]["C2"], book["Room_Y"]["C2"]
    assert formula_cell.data_type == "s" and formula_cell.value == "=1+1"
    assert link_cell.hyperlink is None


def test_students_without_a_section_still_get_an_attendance_row(tmp_path: Path):
    students = pd.DataFrame(
        [["1", "A", 80.0, "CSIT", "A"], ["2", "B", 80.0, "CSIT", None]],
        columns=["Roll Number", "Name", "Attendance Percentage", "Branch", "Section"],
    )
    path = export_sirt_attendance_sheets_excel(students, tmp_path)
    cells = {
        str(c.value)
        for sheet in openpyxl.load_workbook(path).worksheets
        for row in sheet.iter_rows()
        for c in row
    }
    assert {"1", "2"} <= cells


def test_debarred_list_keeps_students_with_a_blank_section(tmp_path: Path):
    debarred = pd.DataFrame(
        [["9", "Z", 10.0, "CSIT", None]],
        columns=["Roll Number", "Name", "Attendance Percentage", "Branch", "Section"],
    )
    path = export_sirt_debarred_excel(debarred, tmp_path)
    book = openpyxl.load_workbook(path)
    assert any("9" in {str(c.value) for row in ws.iter_rows() for c in row} for ws in book.worksheets)


def test_pdf_survives_room_names_with_markup_characters(tmp_path: Path):
    assert export_pdf({"A<B & C": _room("1", "A")}, tmp_path).exists()
